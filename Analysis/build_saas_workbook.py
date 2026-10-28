#!/usr/bin/env python3
"""Build Analysis/output/SaaS_Q1_2026_raw_data.xlsx from the raw bundle files.
Usage: python Analysis/build_saas_workbook.py
"""
from __future__ import annotations

import csv
import io
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

A = Path(__file__).resolve().parent
ROOT = A.parent
sys.path.insert(0, str(A))
from nwlib.sources import load_doc  # noqa: E402

F = "Arial"
H_FONT = Font(name=F, bold=True, color="FFFFFF")
H_FILL = PatternFill("solid", fgColor="1F3864")
BASE = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
BLUE = Font(name=F, size=10, color="0000FF")
GREEN = Font(name=F, size=10, color="008000")
RED_FILL = PatternFill("solid", fgColor="F8CBAD")      # conflict
AMBER_FILL = PatternFill("solid", fgColor="FFE699")    # needs attention / not in scope
GREY_FILL = PatternFill("solid", fgColor="EDEDED")     # subtotal / structure
YELLOW = PatternFill("solid", fgColor="FFFF00")        # key assumption
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))
WRAP = Alignment(wrap_text=True, vertical="top")
USD = '$#,##0;($#,##0);"-"'
AUTHOR = "Analysis"


def note(cell, text, w=360, h=140):
    c = Comment(text, AUTHOR)
    c.width, c.height = w, h
    cell.comment = c


def header(ws, row, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font, cell.fill, cell.alignment = H_FONT, H_FILL, Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


# ---------------------------------------------------------------------- commentary keyed by source line
PNL_COMMENTS = {
    9: ("CONFLICT", "Salesforce: Q1 posted $12,000 (= $48k/yr run-rate). Renewal at $60,000/yr closed 2026-05-14 "
        "(slack_export.md L119; finance_review_2026-05-28.md L23) → $15,000/qtr from Q2. Finance note L23 says posted figure "
        "'looks higher than a $60k annual would imply' — arithmetic says LOWER ($12,000 < $15,000). Renewal post-dates Q1, "
        "so Q1 stays $12,000. See Conflicts C2, C3."),
    12: ("CONFLICT", "Amplitude appears twice (this row + L13), identical amount and description. Maya: '90% sure that's the "
         "same product entered twice' (slack_export.md L147). Priya: 'smells like a double-count' (L150). Invoice not back "
         "(L153). Kept in booked total; excluded in 'likely'. See Conflict C1."),
    13: ("CONFLICT", "Second Amplitude line — suspected duplicate of L12 (slack_export.md L147, L150, L153). Unconfirmed: "
         "invoice not in bundle. This is the $7,500 between booked ($81,000) and likely ($73,500)."),
    24: ("CHECK", "Listed subtotal $81,000. Re-added L9–L23 = $81,000 → ties. Matches CFO 'subtotal as listed is $81k' "
         "(slack_export.md L144). Includes the suspected Amplitude duplicate."),
    25: ("OUT OF SCOPE", "AWS is Infrastructure, not SaaS — CFO: 'don't let anyone fold them together' (slack_export.md L144)."),
    26: ("DATA QUALITY", "Raw value is text '€1900' — EUR, not USD, and not numeric. P&L note says normalize before totaling. "
         "Infrastructure, so excluded from SaaS either way. See Conflict C4."),
    27: ("DATA QUALITY", "Infrastructure subtotal is blank — 'pending FX normalization of EUR line'. Not SaaS."),
    35: ("DATA QUALITY", "Office supplies blank — 'pending AP coding'. Not SaaS; shows P&L is not fully closed."),
}

EVIDENCE = [  # (file, line, relevance, conflicts_with)
    ("slack_export.md", 42, "VPs asked to sanity-check P&L line items; Priya owns the master.", ""),
    ("slack_export.md", 45, "CFO wants a second set of eyes on a couple of SaaS line items.", "Signals known issues in SaaS lines"),
    ("slack_export.md", 48, "Eng tools (GitHub, Sentry) unchanged in Q1.", "Consistent with P&L L20–L21"),
    ("slack_export.md", 119, "Salesforce renewal closed: $60k for the year.", "P&L L9 ($12,000/qtr) — C2"),
    ("slack_export.md", 125, "Maya noting renewal for the SaaS tracker.", "SaaS tracker itself not in bundle"),
    ("slack_export.md", 144, "CFO: SaaS finalized at $81k; AWS (~$38k) is Infrastructure, not SaaS.", "P&L L24 (ties); scope rule for L25–L26"),
    ("slack_export.md", 147, "Maya: Amplitude + Amplitude Analytics, each $7.5k — likely same product twice.", "P&L L12–L13 — C1"),
    ("slack_export.md", 150, "Priya: 'smells like a double-count'; pull invoice.", "P&L L12–L13 — C1"),
    ("slack_export.md", 153, "Maya: invoice not back; leaving CSV as-is.", "C1 unresolved at end of bundle"),
    ("meeting_notes/finance_review_2026-05-28.md", 23, "Salesforce $60k/yr; claims posted Q figure looks higher than $60k/yr implies.", "Arithmetic contradicts it — C3"),
    ("meeting_notes/finance_review_2026-05-28.md", 24, "HubSpot, Looker, Amplitude auto-renew within two quarters.", "Context; no amounts"),
    ("meeting_notes/finance_review_2026-05-28.md", 25, "EU-billed tools come in EUR; normalize before SaaS subtotal.", "No SaaS line is marked EUR — C4"),
    ("meeting_notes/finance_review_2026-05-28.md", 42, "Footer: 'Figures illustrative / pre-close'.", "Lowers weight of this note"),
    ("meeting_notes/data_review_2026-05-20.md", 36, "Amplitude *event* duplicates (post_published) fixed.", "NOT the billing duplicate — different issue, don't conflate"),
]

CONFLICTS = [
    ("C1", "Amplitude listed twice ($7,500 each)", "pnl_q1_2026.csv L12, L13",
     "slack_export.md L147, L150, L153", "Unresolved — invoice not received by end of bundle",
     "Booked $81,000 vs likely $73,500 (−$7,500)", "Medium", "Pull Amplitude invoice/contract (owner: Maya)"),
    ("C2", "Salesforce Q1 posting vs renewed contract", "pnl_q1_2026.csv L9 ($12,000)",
     "slack_export.md L119; finance_review L23 ($60k/yr = $15,000/qtr)",
     "Renewal closed 2026-05-14 (after Q1). Q1 posting stands", "None for Q1; +$3,000/qtr run-rate from Q2", "High",
     "Confirm Q2 posting = $15,000"),
    ("C3", "Finance note says Salesforce posting looks too HIGH", "finance_review_2026-05-28.md L23",
     "Arithmetic: $12,000 < $15,000", "Note is internally inconsistent; not relied on", "None", "High",
     "Ask Maya what she compared against"),
    ("C4", "EUR-billed tools must be normalized — but no SaaS line is in EUR", "finance_review_2026-05-28.md L25",
     "pnl_q1_2026.csv L9–L23 (all plain numbers); only EUR line is Cloudflare L26 (Infrastructure)",
     "Assume SaaS lines already USD", "Unknown; would move total by any FX difference", "Medium",
     "Confirm which tools are EU-billed"),
    ("C5", "SaaS scope: AWS / Cloudflare", "pnl_q1_2026.csv L25–L27 (Infrastructure)", "slack_export.md L144 (CFO)",
     "Excluded from SaaS — sources agree", "None", "High", "—"),
]


def main() -> Path:
    wb = Workbook()

    # ------------------------------------------------------------------ Sheet 2: raw P&L
    raw = wb.create_sheet("P&L raw")
    doc = load_doc(ROOT, ROOT / "pnl_q1_2026.csv")
    raw["A1"] = f"Raw export: {doc.rel} — every row as delivered (values unchanged; € and blanks kept as-is)."
    raw["A1"].font = BOLD
    cols = ["Source line", "Category", "Line Item", "Q1_2026_USD (raw)", "Notes (raw)", "Row type", "In SaaS scope?",
            "Flag", "Commentary & references"]
    header(raw, 2, cols, [9, 18, 26, 14, 40, 12, 10, 13, 90])
    hdr = next(csv.reader(io.StringIO(doc.line(1))))
    raw["A3"] = 1
    for j, v in enumerate(hdr, 2):
        raw.cell(row=3, column=j, value=v).font = Font(name=F, size=10, italic=True)
    raw["F3"] = "header"
    r = 4
    row_of = {}
    for n in range(2, len(doc.lines) + 1):
        text = doc.line(n)
        if not text.strip():
            continue
        c = next(csv.reader(io.StringIO(text))) + ["", "", "", ""]
        cat, item, amt, notes = (x.strip() for x in c[:4])
        try:
            val = int(amt) if amt else None
        except ValueError:
            val = amt                              # keep raw text e.g. '€1900'
        rtype = "section" if not item else ("subtotal" if "total" in item.lower() else "line item")
        in_scope = "Yes" if cat == "Software & SaaS" and rtype == "line item" else ("Subtotal" if cat == "Software & SaaS" else "No")
        vals = [n, cat, item, val, notes, rtype, in_scope]
        for j, v in enumerate(vals, 1):
            cell = raw.cell(row=r, column=j, value=v)
            cell.font, cell.alignment, cell.border = BASE, WRAP, THIN
        raw.cell(row=r, column=4).number_format = USD
        if rtype != "line item":
            for j in range(1, 10):
                raw.cell(row=r, column=j).fill = GREY_FILL
        if n in PNL_COMMENTS:
            flag, txt = PNL_COMMENTS[n]
            raw.cell(row=r, column=8, value=flag).font = BOLD
            raw.cell(row=r, column=9, value=txt).font = BASE
            raw.cell(row=r, column=9).alignment = WRAP
            fill = RED_FILL if flag == "CONFLICT" else AMBER_FILL
            for j in (3, 4, 8):
                raw.cell(row=r, column=j).fill = fill
            note(raw.cell(row=r, column=4), txt)
        row_of[n] = r
        r += 1
    last_raw = r - 1

    # ------------------------------------------------------------------ Sheet 3: evidence
    ev = wb.create_sheet("Evidence (Slack & notes)")
    ev["A1"] = "Every Slack / meeting-note line that bears on SaaS spend — text quoted as written, with exact file and line."
    ev["A1"].font = BOLD
    header(ev, 2, ["File", "Line", "Who", "When", "Raw text", "What it says", "Relates to / conflicts with"],
           [36, 7, 16, 17, 80, 44, 40])
    from nwlib.sources import load_notes, load_slack
    idx = {(m.cite.file, m.cite.line): m for m in load_slack(ROOT) + load_notes(ROOT)}
    for i, (f, ln, what, rel) in enumerate(EVIDENCE, 3):
        m = idx.get((f, ln))
        text = load_doc(ROOT, ROOT / f).line(ln).strip()
        who = m.who if m else ""
        when = m.cite.when if m else ""
        for j, v in enumerate([f, ln, who, when, text, what, rel], 1):
            cell = ev.cell(row=i, column=j, value=v)
            cell.font, cell.alignment, cell.border = BASE, WRAP, THIN
        if "C1" in rel or "C2" in rel or "C3" in rel or "C4" in rel:
            ev.cell(row=i, column=7).fill = RED_FILL
        if "NOT" in rel:
            ev.cell(row=i, column=7).fill = AMBER_FILL

    # ------------------------------------------------------------------ Sheet 4: conflicts
    cf = wb.create_sheet("Conflicts")
    cf["A1"] = "Where sources disagree — what each says, how it was resolved, and what it does to the number."
    cf["A1"].font = BOLD
    header(cf, 2, ["ID", "Issue", "Source A", "Source B", "Resolution", "Effect on Q1 SaaS", "Confidence", "Unblocker"],
           [6, 36, 32, 40, 38, 30, 11, 30])
    for i, row in enumerate(CONFLICTS, 3):
        for j, v in enumerate(row, 1):
            cell = cf.cell(row=i, column=j, value=v)
            cell.font, cell.alignment, cell.border = BASE, WRAP, THIN

    # ------------------------------------------------------------------ Sheet 1: summary (formulas)
    s = wb.active
    s.title = "Summary"
    s.column_dimensions["A"].width = 46
    s.column_dimensions["B"].width = 16
    s.column_dimensions["C"].width = 90
    s["A1"] = "Q1 2026 SaaS spend — Northwind Media"
    s["A1"].font = Font(name=F, size=14, bold=True)
    s["A2"] = ("All figures are live formulas on the 'P&L raw' tab. Blue = typed input (with source). "
               "Red rows on 'P&L raw' = conflicting sources; see 'Conflicts'.")
    s["A2"].font = Font(name=F, size=9, italic=True)
    rng = lambda col: f"'P&L raw'!${col}$4:${col}${last_raw}"  # noqa: E731
    r12, r13, r9, r24 = row_of[12], row_of[13], row_of[9], row_of[24]
    rows = [
        ("Booked SaaS spend (sum of 15 line items)",
         f'=SUMIFS({rng("D")},{rng("B")},"Software & SaaS",{rng("F")},"line item")', USD, BASE,
         "pnl_q1_2026.csv L9–L23. Recomputed from raw rows."),
        ("Listed subtotal in P&L", f"='P&L raw'!D{r24}", USD, GREEN, "pnl_q1_2026.csv L24 'As-listed sum of SaaS line items'."),
        ("Tie-out (booked − listed)", "=B4-B5", USD, BASE, "0 = subtotal is internally consistent. CFO also quotes $81k (slack_export.md L144)."),
        ("Suspected duplicate: 2nd Amplitude line", f"='P&L raw'!D{r13}", USD, GREEN,
         "pnl_q1_2026.csv L13 vs L12 — same amount & description. Flagged by Maya (slack L147) and Priya (L150); "
         "invoice not received (L153). Conflict C1."),
        ("Likely true SaaS spend (booked − duplicate)", "=B4-B7", USD, BOLD, "Medium confidence until the Amplitude invoice is checked."),
        ("", None, None, BASE, ""),
        ("Salesforce check (Conflicts C2, C3)", None, None, BOLD, ""),
        ("Salesforce posted, Q1", f"='P&L raw'!D{r9}", USD, GREEN, "pnl_q1_2026.csv L9."),
        ("Salesforce renewed contract, per year", 60000, USD, BLUE,
         "Input. Source: slack_export.md L119 (Priya, 2026-05-14) and finance_review_2026-05-28.md L23 — '$60k for the year'."),
        ("Implied per quarter under new contract", "=B12/4", USD, BASE, "Applies from renewal (2026-05-14, i.e. Q2) — not to Q1."),
        ("Posted Q1 vs new quarterly rate", "=B11-B13", USD, BASE,
         "Negative = posting is LOWER than the new rate. Contradicts finance note L23 ('looks higher'). Q2+ run-rate rises by this amount."),
    ]
    for i, (label, val, fmt, font, src) in enumerate(rows, 4):
        s.cell(row=i, column=1, value=label).font = BOLD if font is BOLD and val is None else BASE
        if val is not None:
            c = s.cell(row=i, column=2, value=val)
            c.font, c.number_format = font, fmt
        s.cell(row=i, column=3, value=src).font = BASE
        s.cell(row=i, column=3).alignment = WRAP
    s["B8"].fill = YELLOW
    note(s["B7"], "Excluded in 'likely' only. Booked figure keeps it, because it is in the ledger. Resolve with the Amplitude invoice.")
    note(s["B12"], "Hardcoded input — source: slack_export.md L119; meeting_notes/finance_review_2026-05-28.md L23.")
    s["A17"] = "Answer"
    s["A17"].font = BOLD
    s["A18"] = ('="Booked "&TEXT(B4,"$#,##0")&"; likely "&TEXT(B8,"$#,##0")&" (range "&TEXT(B8,"$#,##0")&" – "'
                '&TEXT(B4,"$#,##0")&"). Gap = suspected Amplitude double-count."')
    s["A18"].font = BOLD
    s["A20"] = "Scope & caveats"
    s["A20"].font = BOLD
    caveats = [
        "Scope: Q1 2026 (Jan–Mar), category 'Software & SaaS' only. AWS and Cloudflare are Infrastructure (CFO, slack L144).",
        "Not verifiable from the bundle: vendor invoices; SaaS paid on cards/expenses (would push spend ABOVE booked); "
        "whether any SaaS line is EUR (finance_review L25); whether the P&L is final (notes marked 'pre-close').",
        "Colour key on 'P&L raw': red = conflicting sources · amber = data quality / out of scope · grey = section or subtotal row.",
    ]
    for i, t in enumerate(caveats, 21):
        s.cell(row=i, column=1, value=t).font = BASE
        s.merge_cells(start_row=i, start_column=1, end_row=i, end_column=3)
        s.cell(row=i, column=1).alignment = WRAP
        s.row_dimensions[i].height = 30

    out = A / "output" / "SaaS_Q1_2026_raw_data.xlsx"
    out.parent.mkdir(exist_ok=True)
    wb.save(out)
    print(out)
    return out


if __name__ == "__main__":
    main()
